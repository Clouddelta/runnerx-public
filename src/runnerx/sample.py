"""Generate reproducible, entirely synthetic training-camp input files."""
from __future__ import annotations

import csv
import io
import random
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


START_DATE = date(2026, 1, 5)
END_DATE = date(2026, 2, 1)
HEADERS = ["跑者编号", "姓名", "性别", "出生年份", "身高(cm)", "体重(kg)", "10K测试成绩", "全马最佳成绩", "备赛周期跑量", "经历与计划"]


def _xlsx(path: Path, rows: list[list]) -> None:
    workbook = Workbook()
    workbook.properties.creator = "RunnerX synthetic sample generator"
    workbook.properties.title = "Synthetic training camp registration sample"
    workbook.properties.created = datetime(2026, 1, 1)
    workbook.properties.modified = datetime(2026, 1, 1)
    sheet = workbook.active
    sheet.title = "Synthetic Registrations"
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="173D50")
    for column in ("A", "B", "C", "D", "E", "F", "G", "H", "I", "J"):
        sheet.column_dimensions[column].width = 30 if column in ("B", "J") else 20
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    # Normalize ZIP timestamps and generated document metadata so repeated
    # generation with the same seed produces byte-identical XLSX files.
    with ZipFile(io.BytesIO(buffer.getvalue())) as source, ZipFile(path, "w", compression=ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            content = source.read(name)
            if name == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*(</dcterms:(?:created|modified)>)", rb"\g<1>2026-01-01T00:00:00Z\g<2>", content)
            entry = ZipInfo(name, (2026, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            target.writestr(entry, content)


def generate_sample(output_dir: Path, runners: int = 30, seed: int = 42) -> dict:
    """Write Chinese-header XLSX, four weeks of CSV, and a rejection example.

    No input or identity comes from the legacy spreadsheets. The supplied
    seed controls every generated value and all workbook timestamps are fixed.
    """
    if isinstance(runners, bool) or not isinstance(runners, int) or not 1 <= runners <= 10000:
        raise ValueError("runners must be an integer between 1 and 10000")
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    rng = random.Random(seed)
    registration_rows, session_rows = [], []
    for index in range(1, runners + 1):
        external_id = f"SYN{index:04d}"
        base_pace = rng.randint(285, 410)
        ten_k = rng.randint(2400, 4200)
        marathon = rng.randint(10800, 18000)
        registration_rows.append([
            external_id, f"Sample Runner{index:04d}", "男" if index % 2 else "女",
            rng.randint(1970, 2005), rng.randint(155, 190), round(rng.uniform(48, 90), 1),
            f"{ten_k // 60}:{ten_k % 60:02d}",
            f"{marathon // 3600}小时{marathon % 3600 // 60}分{marathon % 60}秒",
            rng.randint(200, 700), "SYNTHETIC: four-week training sample; no real person.",
        ])
        for week in range(4):
            for day_offset in (0, 2, 5):
                session_date = START_DATE + timedelta(days=week * 7 + day_offset)
                pace = base_pace + rng.randint(-12, 20)
                session_rows.append([
                    external_id, session_date.isoformat(), round(rng.uniform(5, 18), 2),
                    f"{pace // 60}:{pace % 60:02d}", rng.randint(45, 75), rng.randint(1, 5),
                    "SYNTHETIC training record",
                ])
    paths = {
        "registrations": output_dir / "registrations.xlsx",
        "sessions": output_dir / "sessions.csv",
        "invalid_registrations": output_dir / "invalid_registrations.xlsx",
        "readme": output_dir / "README.md",
    }
    _xlsx(paths["registrations"], registration_rows)
    with paths["sessions"].open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream, lineterminator="\n")
        writer.writerow(["external_id", "session_date", "distance_km", "pace_sec_per_km", "resting_hr", "fatigue_level", "notes"])
        writer.writerows(session_rows)
    invalid_rows = [list(registration_rows[0]), list(registration_rows[-1])]
    invalid_rows[0][0] = ""
    invalid_rows[1][0] = "SYN-INVALID"
    invalid_rows[1][6] = "00:99:00"
    _xlsx(paths["invalid_registrations"], invalid_rows)
    paths["readme"].write_text(
        "# Synthetic RunnerX sample data\n\n"
        "All names, identifiers, measurements, and training records in this directory are generated. "
        "They do not describe real people and were not copied from any legacy runner spreadsheet.\n\n"
        f"Generation: runners={runners}, seed={seed}. The same options produce byte-identical files.\n\n"
        "Suggested bootcamp: code `demo-2026`, start date `2026-01-05`, end date `2026-02-01`.\n\n"
        f"- `registrations.xlsx`: {runners} synthetic runners, Chinese column aliases, explicit birth years.\n"
        f"- `sessions.csv`: {len(session_rows)} synthetic sessions, three per runner per week for four weeks.\n"
        "- `invalid_registrations.xlsx`: two deliberately invalid rows (missing external ID; malformed 10K clock). "
        "The entire batch should be rejected, with no business rows changed.\n\n"
        "Import registrations before sessions. Runner identity is `external_id` within a tenant; names are not keys. "
        "Repeated committed imports of the same bytes are skipped.\n\n"
        "Numeric race times are seconds. Race times also accept MM:SS, HH:MM:SS, or Chinese time units. "
        "Pace is seconds per kilometre or MM:SS; a race-duration HH:MM:SS value is rejected as pace.\n",
        encoding="utf-8",
    )
    return paths
